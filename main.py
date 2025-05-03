import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import os
from datetime import datetime
import excel_service
import subprocess
import platform
EXCEL_FILE = "purchases.xlsx"

def save_purchase():
    name = entry_name.get()
    amount = entry_amount.get()
    category = entry_category.get()
    
    if not name or not amount or not category:
        messagebox.showwarning("Missing Field", "All fields are required.")
        return

    try:
        amount = float(amount)
    except ValueError:
        messagebox.showerror("Error", "Amount must be a number.")
        return

    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses"
        ws.append(["Date", "Name", "Amount", "Category"])
    else:
        wb = load_workbook(EXCEL_FILE)
        ws = wb["Expenses"]

    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), name, amount, category])
    try:
        wb.save(EXCEL_FILE)
    except PermissionError:
        fallback_file = f"purchases_queue/purchase_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        os.makedirs("purchases_queue", exist_ok=True)
        wb.save(fallback_file)
        messagebox.showinfo("Saved Temporarily", f"Excel is open. Saved to queue:\n{fallback_file}")


    entry_name.delete(0, tk.END)
    entry_amount.delete(0, tk.END)
    entry_category.delete(0, tk.END)

    messagebox.showinfo("Success", "Purchase recorded!")
excel_service.close_excel_file()
root = tk.Tk()
root.title("Purchase Recorder")
root.geometry("300x250")

tk.Label(root, text="Purchase Name").pack(pady=5)
entry_name = tk.Entry(root, width=30)
entry_name.pack()

tk.Label(root, text="Amount ($)").pack(pady=5)
entry_amount = tk.Entry(root, width=30)
entry_amount.pack()

tk.Label(root, text="Category").pack(pady=5)
entry_category = tk.Entry(root, width=30)
entry_category.pack()

tk.Button(root, text="Save Purchase", command=save_purchase).pack(pady=20)

root.mainloop()
excel_service.merge_temp_files()
if platform.system() == "Windows":
    subprocess.Popen(["start", EXCEL_FILE], shell=True)
