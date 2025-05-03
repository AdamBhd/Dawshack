from openpyxl import load_workbook
import os
import win32com.client
import pythoncom
MAIN_FILE = "purchases.xlsx"
QUEUE_FOLDER = "purchases_queue"
def close_excel_file():
    try:
        pythoncom.CoInitialize()
        excel = win32com.client.GetActiveObject("Excel.Application")
        for wb in excel.Workbooks:
            if wb.FullName.lower() == os.path.abspath(MAIN_FILE).lower():
                wb.Close(SaveChanges=True)
                print(f"Closed Excel file: {MAIN_FILE}")
                break
        if excel.Workbooks.Count == 0:
            excel.Quit()
    except Exception as e:
        print("Excel not open or error closing:", e)


def merge_temp_files():
    if not os.path.exists(MAIN_FILE):
        print(f"{MAIN_FILE} does not exist. Please create it first.")
        return

    main_wb = load_workbook(MAIN_FILE)
    main_ws = main_wb["Expenses"]

    merged_files = 0

    for file in os.listdir(QUEUE_FOLDER):
        if file.endswith(".xlsx"):
            path = os.path.join(QUEUE_FOLDER, file)
            temp_wb = load_workbook(path)
            temp_ws = temp_wb.active

            for row in temp_ws.iter_rows(min_row=2, values_only=True):
                main_ws.append(row)

            merged_files += 1
            os.remove(path)

    main_wb.save(MAIN_FILE)
 

