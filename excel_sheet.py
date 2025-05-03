from openpyxl import Workbook
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl import Workbook, load_workbook
import os
from datetime import datetime
import transaction
import const
import win32com.client
import pythoncom
income_col = 'B'
date_col = 'G'
expense_name_col = 'B'
category_col = 'C'
amount_col = 'D'
frequency_col = 'E'
transactionList = transaction.get_transactions(transaction.initialize_transaction())


def close_excel_file():
    try:
        pythoncom.CoInitialize()
        excel = win32com.client.GetActiveObject("Excel.Application")
        for wb in excel.Workbooks:
            if wb.FullName.lower() == os.path.abspath(const.FILE_NAME).lower():
                wb.Close(SaveChanges=True)
                print(f"Closed Excel file: {const.FILE_NAME}")
                break
        if excel.Workbooks.Count == 0:
            excel.Quit()
    except Exception as e:
        print("Excel not open or error closing:", e)

def create_finance_sheet():
    return load_workbook(const.FINANCE_SHEET_NAME)
    
def find_next_empty_row(ws, column):
    row = 8  
    while ws[f'{column}{row}'].value is not None:
        row += 1
    return row

def create_sheet(wb, date):
    if date in wb.sheetnames:
        return wb[date]
    active_sheet = wb.active
    active_sheet.title = date
    return active_sheet

def add_income(ws, amount):
    ws[f'{income_col}{3}'] += amount

def add_expense(ws, amount, expense_name, category, date):
    next_row = find_next_empty_row(ws, income_col)
    ws[f'{expense_name_col}{next_row}'] = expense_name
    ws[f'{category_col}{next_row}'] = category
    ws[f'{amount_col}{next_row}'] = amount
    ws[f'{date_col}{next_row}'] = date
    ws[f'{frequency_col}{next_row}'] = "Monthly"

def add_bank_transaction(wb):

    for transaction in transactionList:
        date_str = transaction['date'].strftime('%Y-%m-%d') 
        category = transaction['category'][0] if transaction['category'] and len(transaction['category']) > 0 else "Other"
        ws = create_sheet(wb, date_str)
        add_expense(ws, transaction['amount'], transaction['name'], category, date_str)
        


